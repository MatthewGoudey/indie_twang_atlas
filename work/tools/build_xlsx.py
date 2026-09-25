#!/usr/bin/env python3
"""Build V_Album_Atlas.xlsx from work/master.jsonl and the supporting JSON files.
Usage: python3 build_xlsx.py [out.xlsx]
Optional inputs (sheets degrade gracefully if missing):
  work/sheets/scenes.json    [{scene, place, years, key_labels, key_artists, note}]
  work/sheets/labels.json    [{label, base, years_active, zones_lanes, why}]
  work/sheets/artists.json   {artist: {active_years, related, summary}}
  work/sheets/paths.json     [{path, steps:[{id, connection}]}]
  work/borderline_log.jsonl  {"item","decision","reason"}
  work/qa_log.jsonl          {"run","phase","scope","check","result","action"}
"""
import json, os, re, sys, collections, unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

W = "/home/claude/work"
OUT = sys.argv[1] if len(sys.argv) > 1 else f"{W}/deliverables/V_Album_Atlas.xlsx"
S = f"{W}/sheets"

def jl(p):
    return [json.loads(l) for l in open(p, encoding="utf-8") if l.strip()] if os.path.exists(p) else []
def js(p, d):
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else d
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower().replace("&", "and")
    s = re.sub(r"^the\s+", "", s); return re.sub(r"[^a-z0-9]", "", s)

lanes = json.load(open(f"{W}/skeleton/lanes.json"))
LANE = {l["id"]: l for l in lanes}
vocab = json.load(open(f"{W}/skeleton/vocab.json"))
tags = json.load(open(f"{W}/skeleton/tags.json"))
rows = jl(f"{W}/master.jsonl")
rows.sort(key=lambda r: r["id"])
nv = jl(f"{W}/needs_verification.jsonl")
scenes = js(f"{S}/scenes.json", [])
labels_x = js(f"{S}/labels.json", [])
artists_x = js(f"{S}/artists.json", {})
paths = js(f"{S}/paths.json", [])
blog = jl(f"{W}/borderline_log.jsonl")
qalog = jl(f"{W}/qa_log.jsonl")

# ---- field guide
fg = load_workbook(f"{W}/sources/Slacker_Rock_Alt_Country_Field_Guide.xlsx", read_only=True)
fg_art = [r for r in fg["Artists"].iter_rows(min_row=2, values_only=True) if r[1]]
fg_labels = [r for r in fg["Labels"].iter_rows(min_row=2, values_only=True) if r[0]]
FG_ALIAS = {
    "Neil Young & Crazy Horse": ["Neil Young", "Neil Young & Crazy Horse", "Crazy Horse", "Neil Young and Crazy Horse"],
    'Palace / Bonnie "Prince" Billy': ["Palace Brothers", "Palace Music", "Palace Songs", "Palace", "Bonnie 'Prince' Billy", 'Bonnie "Prince" Billy', "Will Oldham"],
    "Courtney Barnett (recent)": ["Courtney Barnett"],
}
fg_norm = set()
for r in fg_art:
    name = r[1]
    for part in FG_ALIAS.get(name, [p.strip() for p in name.split("/")]):
        fg_norm.add(norm(part))
def in_fg(a): return "Y" if norm(a) in fg_norm else ""

# ---- styles
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FT = Font(name="Arial", size=14, bold=True)
HEAD = PatternFill("solid", fgColor="3B4A5A")
ZFILL = {"Core": "FCE9D9", "V": "E3EEF8", "Context": "EDEDED"}
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

wb = Workbook()
def sheet(title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i); c.font = FB; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"
    return ws
def finish(ws, wrap_cols=()):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F
            c.alignment = WRAP if c.column in wrap_cols else TOP
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row,2)}"

lane_label = lambda lid: f"{lid} {LANE[lid]['name']}" if lid in LANE else lid

# ================= README
ws = sheet("README", ["The V — Album Atlas"], [120], first=True)
ws["A1"].font = FT; ws["A1"].fill = PatternFill(None)
readme = [
    "An album-level atlas of indie twang, slacker rock, and everything that fed them. Companion to V_Map.docx and to the Slacker Rock & Alt-Country Field Guide (whose Region and Style vocabulary it reuses).",
    "",
    "THE SHAPE. Picture a V. The bottom point is the modern indie-twang and slacker scene (2014–now). Climbing the V goes back in time through the genres that fed it; the V widens with each era.",
    "LAYER (vertical, from the original release year): L1 Now 2014–2026 · L2 Bridge 1998–2013 · L3 Nineties 1987–1997 · L4 Underground 1978–1986 · L5 Classic 1965–1977 · L6 Roots before 1965 (the R1 roots lane is always L6).",
    "ZONE (horizontal, from the primary lane): Core = the scenes happening now (completionist) · V = direct ancestors and close siblings (deep) · Context = popular branches whose main line leads elsewhere (landmarks only). Row colour: orange = Core, blue = V, grey = Context.",
    "",
    "SHEETS. Albums — one row per album (Layer and Zone are formulas). Artists — one row per artist with key album IDs. Lanes — the 54 lanes with parents/children and a live album count. Scenes, Labels — gazetteer. Tags — the controlled descriptor vocabulary. Paths — the listening paths from V_Map.docx. Borderline Log — every judgment call. Needs Verification — candidates that could not be verified (never in Albums). QA Log — every error-correction run. Stats — live counts by layer × zone, lane, year and priority.",
    "",
    "HOW TO USE. Filter Albums by Zone/Lane/Priority, or search a descriptor (e.g. 'pedal steel'). 'Start here' marks one or two entry points per lane. The Terry column links Josh Terry's No Expectations coverage. Lineage basis: Documented = established by interview, cover, credit, shared member or critic; Inferred = resemblance only.",
    "YOUR COLUMNS. Heard (Y/N), Rating (1–5) and My notes are blank for you; they have dropdowns. Example: Heard = Y, Rating = 4, My notes = 'side two is the one'.",
    "",
    "CONFIDENCE. High = two independent sources agree on artist, title, original year and original label; Medium = one reliable source. Sources lists the verification URLs.",
    "Active years on the Artists sheet are the span of that artist's releases in this atlas, not a full career span.",
]
for t in readme: ws.append([t])
for r in ws.iter_rows(min_row=2):
    for c in r: c.font = F; c.alignment = Alignment(wrap_text=True, vertical="top")

# ================= Albums
AH = ["ID", "Artist", "Album", "Year", "Type", "Label", "Layer", "Zone", "Primary lane", "Secondary lanes", "Style", "Region",
      "Base", "Scene", "Descriptors", "Description", "Lineage", "Lineage basis", "Key tracks", "Priority", "Start here",
      "Terry", "In field guide", "Sources", "Confidence", "Heard", "Rating", "My notes"]
AW = [8, 24, 30, 7, 11, 20, 7, 9, 30, 14, 14, 20, 18, 16, 30, 70, 50, 11, 30, 12, 8, 30, 9, 40, 10, 7, 7, 30]
ws = sheet("Albums", AH, AW)
NA = len(rows)
for i, r in enumerate(rows, 2):
    layer_f = (f'=IF(LEFT(I{i},3)="R1 ","L6",IF(D{i}>=2014,"L1",IF(D{i}>=1998,"L2",IF(D{i}>=1987,"L3",'
               f'IF(D{i}>=1978,"L4",IF(D{i}>=1965,"L5","L6"))))))')
    zone_f = f'=IFERROR(INDEX(Lanes!$C$2:$C$60,MATCH(LEFT(I{i},FIND(" ",I{i})-1),Lanes!$A$2:$A$60,0)),"?")'
    ws.append([r["id"], r["artist"], r["album"], r["year"], r["type"], r["label"], layer_f, zone_f,
               lane_label(r["primary_lane"]), "; ".join(r.get("secondary_lanes", [])), r["style"], r["region"],
               r.get("base", ""), r.get("scene", ""), "; ".join(r["descriptors"]), r["description"], r["lineage"],
               r["lineage_basis"], "; ".join(r["key_tracks"]), r["priority"], r.get("start_here", ""),
               r.get("terry", ""), in_fg(r["artist"]), "\n".join(r["sources"]), r["confidence"], "", "", ""])
finish(ws, wrap_cols=(16, 17, 24))
for i in range(2, NA + 2):
    ws.row_dimensions[i].height = 75
last = NA + 1
for z, col in ZFILL.items():
    ws.conditional_formatting.add(f"A2:AB{last}", FormulaRule(formula=[f'$H2="{z}"'], fill=PatternFill("solid", fgColor=col, bgColor=col)))

# ================= Lanes
kids = collections.defaultdict(list)
for l in lanes:
    for p in l["parents"]:
        if p in LANE and l["id"] not in kids[p]: kids[p].append(l["id"])
by_lane = collections.defaultdict(list)
for r in rows: by_lane[r["primary_lane"]].append(r)
ws = sheet("Lanes", ["ID", "Lane", "Zone", "Era span", "Definition", "Parent lanes", "Child lanes", "Key scenes", "Key labels", "Albums (primary)", "Essential", "Start here"],
           [6, 40, 9, 12, 70, 16, 16, 30, 36, 10, 10, 10])
for l in lanes:
    rs = by_lane[l["id"]]
    sc = [s for s, _ in collections.Counter(r["scene"] for r in rs if r.get("scene")).most_common(4)]
    lb = [s for s, _ in collections.Counter(r["label"] for r in rs if r["label"] not in ("Self-released",)).most_common(5)]
    ch = sorted(set(l.get("children", [])) | set(kids[l["id"]]))
    n = ws.max_row + 1
    ws.append([l["id"], l["name"], l["zone"], l["era"], l["definition"], ", ".join(l["parents"]), ", ".join(ch),
               "; ".join(sc), "; ".join(lb),
               f'=COUNTIF(Albums!$I$2:$I${last},A{n}&" *")',
               f'=COUNTIFS(Albums!$I$2:$I${last},A{n}&" *",Albums!$T$2:$T${last},"Essential")',
               f'=COUNTIFS(Albums!$I$2:$I${last},A{n}&" *",Albums!$U$2:$U${last},"Y")'])
finish(ws, wrap_cols=(5, 8, 9))
NL = len(lanes) + 1

# ================= Artists
arts = collections.OrderedDict()
for r in sorted(rows, key=lambda r: (norm(r["artist"]), r["year"])):
    arts.setdefault(r["artist"], []).append(r)
ws = sheet("Artists", ["Artist", "Base", "Region", "Active years (in atlas)", "Primary lane", "Zone", "Key album IDs", "Albums in atlas",
                       "Related artists", "Summary", "In field guide"], [28, 18, 20, 12, 34, 9, 30, 9, 40, 70, 9])
for a, rs in sorted(arts.items(), key=lambda kv: norm(kv[0])):
    x = artists_x.get(a, {})
    pl = collections.Counter(r["primary_lane"] for r in rs).most_common(1)[0][0]
    pri = {"Essential": 0, "Recommended": 1, "Deep cut": 2}
    key = [r["id"] for r in sorted(rs, key=lambda r: (pri.get(r["priority"], 3), r["year"]))][:4]
    yrs = sorted(r["year"] for r in rs)
    base = collections.Counter(r.get("base", "") for r in rs).most_common(1)[0][0]
    reg = collections.Counter(r["region"] for r in rs).most_common(1)[0][0]
    n = ws.max_row + 1
    ws.append([a, base, reg, x.get("active_years") or (f"{yrs[0]}–{yrs[-1]}" if yrs[0] != yrs[-1] else str(yrs[0])),
               lane_label(pl), LANE[pl]["zone"], ", ".join(key), f'=COUNTIF(Albums!$B$2:$B${last},A{n})',
               x.get("related", ""), x.get("summary", ""), in_fg(a)])
finish(ws, wrap_cols=(9, 10))

# ================= Scenes
ws = sheet("Scenes", ["Scene", "Place", "Years", "Key labels", "Key artists", "Note", "Albums tagged"], [26, 22, 12, 30, 44, 60, 10])
scene_names = [s["scene"] for s in scenes]
for s in scenes:
    n = ws.max_row + 1
    ws.append([s["scene"], s.get("place", ""), s.get("years", ""), s.get("key_labels", ""), s.get("key_artists", ""), s.get("note", ""),
               f'=COUNTIF(Albums!$N$2:$N${last},A{n})'])
finish(ws, wrap_cols=(4, 5, 6))
NS = max(len(scenes) + 1, 2)

# ================= Labels
ws = sheet("Labels", ["Label", "Base", "Years active", "Zones and lanes", "Why it matters", "In field guide", "Albums in atlas"], [28, 20, 12, 30, 70, 9, 10])
seen = set()
lx = {norm(l["label"]): l for l in labels_x}
lab_lanes = collections.defaultdict(collections.Counter)
for r in rows: lab_lanes[norm(r["label"].split("/")[0])][r["primary_lane"]] += 1
def zl(name):
    c = lab_lanes.get(norm(name), {})
    if not c: return ""
    zs = sorted({LANE[k]["zone"] for k in c})
    return "/".join(zs) + ": " + ", ".join(k for k, _ in c.most_common(5))
for r in fg_labels:
    x = lx.get(norm(r[0]), {})
    n = ws.max_row + 1
    ws.append([r[0], r[1], x.get("years_active", ""), x.get("zones_lanes") or zl(r[0]), x.get("why") or r[2], "Y",
               f'=COUNTIF(Albums!$F$2:$F${last},A{n}&"*")'])
    seen.add(norm(r[0]))
for l in labels_x:
    if norm(l["label"]) in seen: continue
    n = ws.max_row + 1
    ws.append([l["label"], l.get("base", ""), l.get("years_active", ""), l.get("zones_lanes") or zl(l["label"]), l.get("why", ""), "",
               f'=COUNTIF(Albums!$F$2:$F${last},A{n}&"*")'])
finish(ws, wrap_cols=(4, 5))

# ================= Tags
ws = sheet("Tags", ["Category", "Tag", "Definition", "Albums using it"], [16, 24, 70, 10])
for cat, d in tags.items():
    for t, dfn in d.items():
        n = ws.max_row + 1
        ws.append([cat, t, dfn, f'=COUNTIF(Albums!$O$2:$O${last},"*"&B{n}&"*")'])
finish(ws, wrap_cols=(3,))
NT = ws.max_row

# ================= Paths
idx = {r["id"]: r for r in rows}
ws = sheet("Paths", ["Path", "Step", "Album ID", "Artist", "Album", "Year", "Lane", "Connection to the next step"], [30, 6, 9, 24, 30, 7, 30, 80])
for p in paths:
    for k, st in enumerate(p["steps"], 1):
        n = ws.max_row + 1
        ws.append([p["path"], k, st["id"],
                   f'=IFERROR(INDEX(Albums!$B$2:$B${last},MATCH(C{n},Albums!$A$2:$A${last},0)),"?")',
                   f'=IFERROR(INDEX(Albums!$C$2:$C${last},MATCH(C{n},Albums!$A$2:$A${last},0)),"?")',
                   f'=IFERROR(INDEX(Albums!$D$2:$D${last},MATCH(C{n},Albums!$A$2:$A${last},0)),"?")',
                   f'=IFERROR(INDEX(Albums!$I$2:$I${last},MATCH(C{n},Albums!$A$2:$A${last},0)),"?")',
                   st.get("connection", "")])
finish(ws, wrap_cols=(8,))

# ================= Borderline Log
ws = sheet("Borderline Log", ["#", "Item", "Album ID", "Decision", "Reason", "Source"], [5, 40, 9, 30, 80, 16])
k = 0
for b in blog:
    k += 1; ws.append([k, b.get("item", ""), b.get("id", ""), b.get("decision", ""), b.get("reason", ""), "orchestrator"])
for r in rows:
    if r.get("borderline_note"):
        k += 1
        ws.append([k, f'{r["artist"]} – {r["album"]} ({r["year"]})', r["id"], f'Primary {r["primary_lane"]}' + (f'; secondary {", ".join(r["secondary_lanes"])}' if r.get("secondary_lanes") else ""),
                   r["borderline_note"], "lane research"])
finish(ws, wrap_cols=(2, 4, 5))

# ================= Needs Verification
ws = sheet("Needs Verification", ["Artist", "Album", "Year (claimed)", "Lane file", "Why not verified"], [26, 34, 10, 22, 90])
for r in nv:
    ws.append([r.get("artist", ""), r.get("album", ""), r.get("year", "") or "", os.path.basename(r.get("_file", "")).replace("lane_", "").replace(".jsonl", ""), r.get("note", "")])
finish(ws, wrap_cols=(5,))

# ================= QA Log
ws = sheet("QA Log", ["Run", "Phase", "Scope", "Check", "Result", "Action taken"], [8, 16, 30, 30, 60, 60])
for q in qalog:
    ws.append([q.get("run", ""), q.get("phase", ""), q.get("scope", ""), q.get("check", ""), q.get("result", ""), q.get("action", "")])
finish(ws, wrap_cols=(3, 4, 5, 6))

# ================= Stats
ws = sheet("Stats", ["Layer × zone", "Core", "V", "Context", "Total", "", "Floor note"], [22, 10, 10, 10, 10, 4, 50])
LAYERS = [("L1", "Now 2014–2026"), ("L2", "Bridge 1998–2013"), ("L3", "Nineties 1987–1997"), ("L4", "Underground 1978–1986"), ("L5", "Classic 1965–1977"), ("L6", "Roots before 1965")]
floors = {"L3": "V floor 300+", "L4": "V floor 200+", "L5": "V floor 250+", "L6": "R1 roots 25–40"}
for L, nm in LAYERS:
    n = ws.max_row + 1
    ws.append([f"{L} {nm}"] + [f'=COUNTIFS(Albums!$G$2:$G${last},"{L}",Albums!$H$2:$H${last},"{z}")' for z in ("Core", "V", "Context")] + [f"=SUM(B{n}:D{n})", "", floors.get(L, "")])
n = ws.max_row + 1
ws.append(["Total", f"=SUM(B2:B{n-1})", f"=SUM(C2:C{n-1})", f"=SUM(D2:D{n-1})", f"=SUM(E2:E{n-1})", "", "Core floor 500+ (all layers)"])
ws.append(["V zone L1+L2", "", "=C2+C3", "", "", "", "V floor 350+"])
ws.append([])
ws.append(["By priority", "Count"]); hp = ws.max_row
for p in vocab["priority"]:
    n = ws.max_row + 1; ws.append([p, f'=COUNTIF(Albums!$T$2:$T${last},A{n})'])
ws.append([])
ws.append(["By lane", "Albums", "Zone", "Name"]); hl = ws.max_row
for j, l in enumerate(lanes, 2):
    n = ws.max_row + 1
    ws.append([f"=Lanes!A{j}", f"=Lanes!J{j}", f"=Lanes!C{j}", f"=Lanes!B{j}"])
ws.append([])
ws.append(["By year", "Albums"]); hy = ws.max_row
for y in range(min(r["year"] for r in rows), max(r["year"] for r in rows) + 1):
    n = ws.max_row + 1; ws.append([y, f'=COUNTIF(Albums!$D$2:$D${last},A{n})'])
finish(ws)
for rr in (1, hp, hl, hy):
    for c in ws[rr]:
        if c.value is not None: c.font = FB; c.fill = HEAD

# ================= Data validation on Albums
wsA = wb["Albums"]
def dv_list(ws, col, items=None, ref=None, allow_blank=True):
    f = ref if ref else '"' + ",".join(items) + '"'
    dv = DataValidation(type="list", formula1=f, allow_blank=allow_blank, showErrorMessage=True)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{max(last, 2) + 500}")
# hidden list sheet for long vocabularies
lst = wb.create_sheet("Lists")
cols = {"styles": vocab["styles"], "regions": vocab["regions"], "lanes": [lane_label(l["id"]) for l in lanes],
        "scenes": scene_names or [""], "tags": [t for d in tags.values() for t in d]}
for j, (k2, vals) in enumerate(cols.items(), 1):
    lst.cell(1, j, k2)
    for i2, v in enumerate(vals, 2): lst.cell(i2, j, v)
lst.sheet_state = "hidden"
def rng(j, k2): return f"Lists!${get_column_letter(j)}$2:${get_column_letter(j)}${len(cols[k2]) + 1}"
dv_list(wsA, "E", vocab["types"])
dv_list(wsA, "K", ref=rng(1, "styles"))
dv_list(wsA, "L", ref=rng(2, "regions"))
dv_list(wsA, "I", ref=rng(3, "lanes"))
dv_list(wsA, "N", ref=rng(4, "scenes"))
dv_list(wsA, "R", vocab["lineage_basis"])
dv_list(wsA, "T", vocab["priority"])
dv_list(wsA, "U", ["Y"])
dv_list(wsA, "W", ["Y"])
dv_list(wsA, "Y", vocab["confidence"])
dv_list(wsA, "Z", ["Y", "N"])
dv_list(wsA, "AA", ["1", "2", "3", "4", "5"])
dv_list(wb["Artists"], "C", ref=rng(2, "regions"))
dv_list(wb["Lanes"], "C", ["Core", "V", "Context"])
dv_list(wb["Tags"], "A", list(tags.keys()))

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}: {NA} albums, {len(arts)} artists, {len(nv)} nv, {len(scenes)} scenes, {len(paths)} paths")
